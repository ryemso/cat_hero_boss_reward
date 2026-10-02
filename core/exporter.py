from __future__ import annotations
from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from . import db
from .settlement import detail_rows, member_summary, member_item_pivot

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SUB_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="B7B7B7")


def _write_df(ws, df: pd.DataFrame, start_row=1, start_col=1, title=None):
    r = start_row
    if title:
        ws.cell(r, start_col, title).font = Font(bold=True, size=13)
        r += 2
    for j, col in enumerate(df.columns, start_col):
        c = ws.cell(r, j, str(col))
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=THIN)
    for i, row in enumerate(df.itertuples(index=False), r + 1):
        for j, val in enumerate(row, start_col):
            ws.cell(i, j, None if pd.isna(val) else val)
    for j, col in enumerate(df.columns, start_col):
        max_len = max([len(str(col))] + [len(str(v)) for v in df[col].head(300).fillna("").tolist()])
        ws.column_dimensions[get_column_letter(j)].width = min(max(max_len + 2, 10), 28)
    ws.freeze_panes = ws.cell(r + 1, start_col)
    return r


def build_xlsx(date_from=None, date_to=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "최종 정산"
    summary = member_summary(date_from, date_to)
    summary = summary.rename(columns={"member":"사용자","total_value":"보라가치","cash":"현금","after_fee":"수수료 제외"})
    _write_df(ws, summary, title="최종 정산")

    ws2 = wb.create_sheet("개인별 분배")
    pivot = member_item_pivot(date_from, date_to)
    _write_df(ws2, pivot, title="개인별 보상 누적")

    ws3 = wb.create_sheet("보스전 기록")
    raids = pd.DataFrame(db.rows("""
        SELECT r.id, r.raid_date 날짜, r.raid_time 시간, b.name 보스, r.memo 메모,
               GROUP_CONCAT(m.name, ', ') 참여자
        FROM raids r LEFT JOIN bosses b ON b.id=r.boss_id
        LEFT JOIN raid_members x ON x.raid_id=r.id
        LEFT JOIN members m ON m.id=x.member_id
        GROUP BY r.id ORDER BY r.raid_date, r.raid_time, r.id
    """))
    _write_df(ws3, raids, title="보스전 기록")

    ws4 = wb.create_sheet("분배 원장")
    detail = detail_rows(date_from, date_to)
    _write_df(ws4, detail, title="분배 계산 원장")

    ws5 = wb.create_sheet("보상 마스터")
    master = pd.DataFrame(db.rows("""
        SELECT id, name 보상명, settlement_column 정산컬럼, rarity 등급,
               quantity_mode 수량방식, distribution_type 분배방식, unit_value 환산가치, enabled 사용
        FROM reward_master ORDER BY id
    """))
    _write_df(ws5, master, title="보상 마스터")

    ws6 = wb.create_sheet("설정")
    settings = pd.DataFrame(db.rows("SELECT key, value FROM settings ORDER BY key"))
    _write_df(ws6, settings, title="설정")

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
