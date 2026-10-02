import sys,tempfile,unittest,shutil
from pathlib import Path
from io import BytesIO
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from core import db
from core.vision import load_template,match_reward,detect_reward_cards
from core.recognition import recognize_card
from core.settlement import member_summary,member_item_pivot
from core.exporter import build_xlsx
from openpyxl import load_workbook

class Regression(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.old=db.DB_PATH;db.DB_PATH=Path(self.tmp.name)/'settlement.db';db.init_db()
    def tearDown(self):
        db.DB_PATH=self.old;self.tmp.cleanup()
    def test_all_U_templates_and_unicode_relocation(self):
        ts=db.get_templates();self.assertEqual({t['candidate_id'] for t in ts},{f'U{i:02}' for i in range(1,34)})
        for t in ts:
            match=match_reward(load_template(t['template_path']),ts)
            self.assertEqual(match['candidate_id'],t['candidate_id']);self.assertGreater(match['score'],.99)
        path=Path(self.tmp.name)/'한글 폴더'/'보상.png';path.parent.mkdir();shutil.copy(ts[0]['template_path'],path)
        self.assertIsNotNone(load_template(path))
        db.execute("UPDATE reward_templates SET template_path='C:\\Users\\예전\\boss\\data\\starter_templates\\U01.png' WHERE candidate_id='U01'")
        db.init_db();self.assertTrue(Path(db.get_templates()[0]['template_path']).is_file())
    def test_master_edit_persists_and_mapping_is_immutable(self):
        r=db.rows("SELECT * FROM reward_master WHERE name='파템'")[0]
        db.update_reward(r['id'],'파템 변경','파템','BLUE','DEFAULT_ONE','EQUAL',1234,False)
        db.init_db()
        updated=db.rows('SELECT * FROM reward_master WHERE id=?',(r['id'],))[0]
        self.assertEqual(updated['unit_value'],1234);self.assertEqual(updated['name'],'파템 변경');self.assertEqual(updated['enabled'],0)
        self.assertEqual(len(db.get_templates()),33);self.assertEqual(len(db.rows('SELECT * FROM reward_master')),12)
        with self.assertRaises(ValueError):db.update_reward(r['id'],'파템 변경','초템','BLUE','AUTO','EQUAL',1,True)
    def test_participant_lifecycle_and_history(self):
        db.save_member(None,'테스트');mid=db.scalar("SELECT id FROM members WHERE name='테스트'")
        db.save_member(mid,'새 이름',False);db.init_db()
        self.assertEqual(db.scalar('SELECT active FROM members WHERE id=?',(mid,)),0)
        db.delete_member(mid);db.init_db();self.assertIsNone(db.scalar('SELECT id FROM members WHERE id=?',(mid,)))
        original=db.scalar("SELECT id FROM members WHERE name='인솔'");db.delete_member(original);db.init_db()
        self.assertIsNone(db.scalar("SELECT id FROM members WHERE name='인솔'"))
    def test_payload_from_real_images(self):
        ts=db.get_templates()
        for path in (ROOT/'tests/fixtures').glob('*.jpg'):
            cards=detect_reward_cards(load_template(path));self.assertGreater(len(cards),0)
            for card in cards:
                row=recognize_card(card,ts,.78,quantity_reader=lambda *a,**kw:(120,.984))
                self.assertNotEqual(row['보상'],'미지정');self.assertTrue(row['후보ID']);self.assertGreater(row['인식점수'],.9)
                if row['OCR수량'] is not None:self.assertEqual(row['수량신뢰도'],.984)
                else:self.assertEqual(row['수량'],1);self.assertIsNone(row['수량신뢰도'])
    def test_missing_and_low_scoring_templates_keep_diagnostics(self):
        cards=detect_reward_cards(load_template(ROOT/'tests/fixtures/KakaoTalk_20261002_135728198.jpg'))
        row=recognize_card(cards[0],[],.78,quantity_reader=lambda *a,**k:(None,0))
        self.assertIsNone(row['인식점수']);self.assertEqual(row['보상'],'미지정');self.assertIn('OCR 실패',row['수량상태'])
        with patch('core.recognition.match_reward',return_value={'id':1,'name':'룬 소환','candidate_id':'U14','quantity_mode':'OCR','score':.6,'accepted':False}):
            row=recognize_card(cards[0],db.get_templates(),.78,quantity_reader=lambda *a,**k:(120,.99))
        self.assertEqual(row['후보ID'],'U14');self.assertEqual(row['인식점수'],.6);self.assertEqual(row['수량'],120)
    def test_settlement_exclusion_export_and_legacy_migration(self):
        ids=[r['id'] for r in db.rows('SELECT id FROM members LIMIT 2')]
        r=db.rows("SELECT * FROM reward_master WHERE name='파템'")[0]
        excluded=db.rows("SELECT * FROM reward_master WHERE name='빨레'")[0]
        db.execute('UPDATE reward_master SET unit_value=999999 WHERE id=?',(excluded['id'],))
        raid=db.save_raid('2026-10-02','14:00',1,ids,[{'reward_id':r['id'],'reward_name':'파템','quantity':2,'score':.987,'candidate_id':'U05','quantity_confidence':.9,'rarity':'BLUE'}, {'reward_id':excluded['id'],'reward_name':'빨레','quantity':1}])
        s=member_summary();self.assertEqual(s.total_value.tolist(),[4300,4300]);self.assertEqual(s.after_fee.tolist(),[19350,19350])
        self.assertEqual(member_item_pivot()['파템'].tolist(),[1,1])
        db.save_member(ids[0],'이름 수정',False);self.assertIn('이름 수정',member_summary().member.tolist())
        with self.assertRaises(ValueError):db.delete_member(ids[0])
        saved=db.rows('SELECT * FROM raid_rewards WHERE raid_id=?',(raid,))[0];self.assertEqual(saved['candidate_id'],'U05');self.assertEqual(saved['quantity_confidence'],.9)
        wb=load_workbook(BytesIO(build_xlsx()));self.assertIn('개인별 분배',wb.sheetnames)
        db.delete_raid(raid);self.assertTrue(member_summary().empty)
    def test_upgrade_existing_database_preserves_records_and_values(self):
        db.DB_PATH=Path(self.tmp.name)/'legacy.db'
        with db.connect() as con:
            con.executescript(db.SCHEMA)
            con.execute("INSERT INTO members(id,name) VALUES(1,'기존 참여자')")
            con.execute("INSERT INTO reward_master(id,name,settlement_column,quantity_mode,unit_value,enabled,created_at) VALUES(1,'파템','파템','DEFAULT_ONE',777,0,'2026-10-01')")
            con.execute("INSERT INTO reward_templates(reward_id,candidate_id,template_path,created_at) VALUES(1,'U05','C:/old/data/starter_templates/U05.png','2026-10-01')")
            con.execute("INSERT INTO raids(id,raid_date,raid_time,created_at) VALUES(1,'2026-10-01','12:00','2026-10-01')")
            con.execute('INSERT INTO raid_members VALUES(1,1)')
            con.execute("INSERT INTO raid_rewards(raid_id,reward_id,reward_name_snapshot,quantity) VALUES(1,1,'파템',2)")
        db.init_db()
        self.assertEqual(db.scalar('SELECT COUNT(*) FROM raid_rewards'),1)
        self.assertEqual(db.scalar('SELECT unit_value FROM reward_master WHERE id=1'),777)
        self.assertEqual(db.scalar('SELECT enabled FROM reward_master WHERE id=1'),0)
        self.assertEqual(member_summary().total_value.tolist(),[1554])
        cols={r['name'] for r in db.rows('PRAGMA table_info(raid_rewards)')}
        self.assertTrue({'candidate_id','quantity_confidence','rarity'} <= cols)

    def test_ui_startup_and_member_edit(self):
        from streamlit.testing.v1 import AppTest
        at=AppTest.from_file(str(ROOT/'app.py')).run(timeout=30)
        self.assertFalse(at.exception)
        new=next(t for t in at.text_input if t.label=='새 참여자');new.set_value('UI 추가')
        next(b for b in at.button if b.label=='참여자 추가').click();at.run()
        self.assertFalse(at.exception);self.assertEqual(db.scalar("SELECT COUNT(*) FROM members WHERE name='UI 추가'"),1)

if __name__=='__main__':unittest.main()
